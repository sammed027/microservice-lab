import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_full_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette matching friend's PPT exact theme
    NAVY_DARK = RGBColor(23, 35, 61)       # #17233d
    NAVY_HEADER = RGBColor(23, 35, 61)     # Header banner
    BG_LIGHT = RGBColor(244, 246, 249)     # Slide body bg #f4f6f9
    TEXT_WHITE = RGBColor(255, 255, 255)
    TEXT_DARK = RGBColor(31, 41, 55)       # Body text #1f2937
    TEXT_MUTED = RGBColor(100, 116, 139)   # Subtitle text #64748b
    TEAL_ACCENT = RGBColor(0, 150, 136)    # Teal border #009688
    ORANGE_ACCENT = RGBColor(230, 126, 34) # Orange box/bar #e67e22
    CARD_BG = RGBColor(255, 255, 255)      # White card bg

    def set_slide_background(slide, color):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()

    def add_top_banner(slide, title_text):
        banner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(1.1))
        banner.fill.solid()
        banner.fill.fore_color.rgb = NAVY_HEADER
        banner.line.fill.background()

        rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(1.1), Inches(13.333), Inches(0.06))
        rule.fill.solid()
        rule.fill.fore_color.rgb = TEAL_ACCENT
        rule.line.fill.background()

        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.25), Inches(11.7), Inches(0.7))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE

    def add_footer(slide, page_num):
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(7.0), Inches(11.7), Inches(0.4))
        tf = tb.text_frame
        p = tf.paragraphs[0]
        p.text = f"Gym Management Microservices | Cloud Computing Lab Evaluation"
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_MUTED

    # -------------------------------------------------------------
    # SLIDE 1: TITLE SLIDE
    # -------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1, NAVY_DARK)

    orange_bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(2.1), Inches(0.15), Inches(2.6))
    orange_bar.fill.solid()
    orange_bar.fill.fore_color.rgb = ORANGE_ACCENT
    orange_bar.line.fill.background()

    tb1 = slide1.shapes.add_textbox(Inches(1.2), Inches(2.0), Inches(11.3), Inches(3.2))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "Build, Deploy and Analyze a Containerized\nMicroservice Application Under Varying Workloads"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE

    p2 = tf1.add_paragraph()
    p2.text = "Gym Management Microservices"
    p2.font.size = Pt(22)
    p2.font.color.rgb = RGBColor(147, 197, 253)
    p2.space_before = Pt(20)

    p3 = tf1.add_paragraph()
    p3.text = "Cloud Computing Lab Evaluation"
    p3.font.size = Pt(18)
    p3.font.color.rgb = RGBColor(203, 213, 225)
    p3.space_before = Pt(8)

    tb_pres = slide1.shapes.add_textbox(Inches(1.2), Inches(6.2), Inches(11.0), Inches(0.6))
    tf_pres = tb_pres.text_frame
    p_pr = tf_pres.paragraphs[0]
    p_pr.text = "Presented by: Student Team   |   KLE Technological University, BVB Campus, Hubli"
    p_pr.font.size = Pt(12)
    p_pr.font.color.rgb = RGBColor(203, 213, 225)

    # -------------------------------------------------------------
    # SLIDE 2: AIM AND OBJECTIVES
    # -------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2, BG_LIGHT)
    add_top_banner(slide2, "Aim and Objectives")

    tb_aim = slide2.shapes.add_textbox(Inches(0.8), Inches(1.35), Inches(11.7), Inches(0.8))
    tf_aim = tb_aim.text_frame
    tf_aim.word_wrap = True
    p = tf_aim.paragraphs[0]
    p.text = "To build a microservice application with three independent services, containerize and deploy it with Docker and Docker Compose, generate varying workloads, monitor resource use, and analyze performance."
    p.font.size = Pt(15)
    p.font.color.rgb = TEXT_DARK

    card_data = [
        ("Develop", "Three independent FastAPI microservices with REST APIs"),
        ("Containerize", "One Dockerfile and one Docker image per service"),
        ("Deploy", "Docker Compose runs all three containers on one network"),
        ("Analyze", "Load test at 1, 2, 4, 8, 16 concurrent requests and monitor CPU and memory")
    ]
    card_width, card_height, start_left, spacing = Inches(2.7), Inches(3.2), Inches(0.8), Inches(0.3)
    for i, (title, desc) in enumerate(card_data):
        left_pos = start_left + i * (card_width + spacing)
        top_pos = Inches(2.4)
        card = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, left_pos, top_pos, card_width, card_height)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = RGBColor(226, 232, 240)

        t_bar = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, left_pos, top_pos, Inches(0.12), card_height)
        t_bar.fill.solid()
        t_bar.fill.fore_color.rgb = TEAL_ACCENT
        t_bar.line.fill.background()

        tb_c = slide2.shapes.add_textbox(left_pos + Inches(0.25), top_pos + Inches(0.2), card_width - Inches(0.3), card_height - Inches(0.4))
        tf_c = tb_c.text_frame
        tf_c.word_wrap = True
        p_t = tf_c.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(18)
        p_t.font.bold = True
        p_t.font.color.rgb = NAVY_DARK

        p_d = tf_c.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(13)
        p_d.font.color.rgb = TEXT_DARK
        p_d.space_before = Pt(12)

    tb_note = slide2.shapes.add_textbox(Inches(0.8), Inches(6.1), Inches(11.7), Inches(0.5))
    p_n = tb_note.text_frame.paragraphs[0]
    p_n.text = "Core challenge: moving from one monolith to independent, networked containers while finding the performance bottleneck."
    p_n.font.size = Pt(13)
    p_n.font.italic = True
    p_n.font.color.rgb = TEAL_ACCENT
    add_footer(slide2, 2)

    # -------------------------------------------------------------
    # SLIDE 3: TECH STACK AND ENVIRONMENT
    # -------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3, BG_LIGHT)
    add_top_banner(slide3, "Technology Stack and Environment")

    tech_grid = [
        ("Backend", "Python 3.11 + FastAPI"),
        ("Web server", "Uvicorn (ASGI worker)"),
        ("Containerization", "Docker 29.8.0"),
        ("Orchestration", "Docker Compose v5.5.1"),
        ("Load testing", "Python thread pool"),
        ("Test machine", "macOS, 6 CPUs, 3.8 GB for Docker")
    ]
    box_w, box_h, col_s, row_s = Inches(3.7), Inches(2.2), Inches(0.3), Inches(0.3)
    for i, (title, desc) in enumerate(tech_grid):
        col, row = i % 3, i // 3
        left_pos = Inches(0.8) + col * (box_w + col_s)
        top_pos = Inches(1.6) + row * (box_h + row_s)
        box = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, left_pos, top_pos, box_w, box_h)
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = RGBColor(226, 232, 240)

        t_bar = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, left_pos, top_pos, Inches(0.12), box_h)
        t_bar.fill.solid()
        t_bar.fill.fore_color.rgb = TEAL_ACCENT
        t_bar.line.fill.background()

        tb_box = slide3.shapes.add_textbox(left_pos + Inches(0.25), top_pos + Inches(0.2), box_w - Inches(0.3), box_h - Inches(0.4))
        tf_b = tb_box.text_frame
        tf_b.word_wrap = True
        p_t = tf_b.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = NAVY_DARK

        p_d = tf_b.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(18)
        p_d.font.bold = True
        p_d.font.color.rgb = TEXT_DARK
        p_d.space_before = Pt(14)
    add_footer(slide3, 3)

    # -------------------------------------------------------------
    # SLIDE 4: SYSTEM ARCHITECTURE
    # -------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4, BG_LIGHT)
    add_top_banner(slide4, "System Architecture")

    client_box = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.2), Inches(1.5), Inches(2.9), Inches(0.8))
    client_box.fill.solid()
    client_box.fill.fore_color.rgb = RGBColor(100, 116, 139)
    client_box.line.fill.background()
    tf_c = client_box.text_frame
    p_c1 = tf_c.paragraphs[0]
    p_c1.text = "Client"
    p_c1.font.size = Pt(16)
    p_c1.font.bold = True
    p_c1.font.color.rgb = TEXT_WHITE
    p_c1.alignment = PP_ALIGN.CENTER
    p_c2 = tf_c.add_paragraph()
    p_c2.text = "curl / load test"
    p_c2.font.size = Pt(11)
    p_c2.font.color.rgb = RGBColor(226, 232, 240)
    p_c2.alignment = PP_ALIGN.CENTER

    net_box = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2.5), Inches(2.7), Inches(8.3), Inches(4.0))
    net_box.fill.solid()
    net_box.fill.fore_color.rgb = RGBColor(240, 253, 250)
    net_box.line.color.rgb = TEAL_ACCENT
    net_box.line.width = Pt(2)

    tb_net = slide4.shapes.add_textbox(Inches(2.7), Inches(2.8), Inches(5.0), Inches(0.4))
    p_n = tb_net.text_frame.paragraphs[0]
    p_n.text = "Docker network: gym-network"
    p_n.font.size = Pt(14)
    p_n.font.bold = True
    p_n.font.color.rgb = TEAL_ACCENT

    att_box = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.5), Inches(3.5), Inches(4.3), Inches(1.1))
    att_box.fill.solid()
    att_box.fill.fore_color.rgb = NAVY_DARK
    att_box.line.fill.background()
    tf_att = att_box.text_frame
    p_a1 = tf_att.paragraphs[0]
    p_a1.text = "attendance-service"
    p_a1.font.size = Pt(18)
    p_a1.font.bold = True
    p_a1.font.color.rgb = TEXT_WHITE
    p_a1.alignment = PP_ALIGN.CENTER
    p_a2 = tf_att.add_paragraph()
    p_a2.text = "Port 8003 (entry point)"
    p_a2.font.size = Pt(12)
    p_a2.font.color.rgb = RGBColor(203, 213, 225)
    p_a2.alignment = PP_ALIGN.CENTER

    mem_box = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.0), Inches(5.3), Inches(3.4), Inches(1.1))
    mem_box.fill.solid()
    mem_box.fill.fore_color.rgb = TEAL_ACCENT
    mem_box.line.fill.background()
    tf_m = mem_box.text_frame
    p_m1 = tf_m.paragraphs[0]
    p_m1.text = "member-service"
    p_m1.font.size = Pt(17)
    p_m1.font.bold = True
    p_m1.font.color.rgb = TEXT_WHITE
    p_m1.alignment = PP_ALIGN.CENTER
    p_m2 = tf_m.add_paragraph()
    p_m2.text = "Port 8001"
    p_m2.font.size = Pt(12)
    p_m2.font.color.rgb = TEXT_WHITE
    p_m2.alignment = PP_ALIGN.CENTER

    ms_box = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(5.3), Inches(3.4), Inches(1.1))
    ms_box.fill.solid()
    ms_box.fill.fore_color.rgb = ORANGE_ACCENT
    ms_box.line.fill.background()
    tf_ms = ms_box.text_frame
    p_ms1 = tf_ms.paragraphs[0]
    p_ms1.text = "membership-service"
    p_ms1.font.size = Pt(17)
    p_ms1.font.bold = True
    p_ms1.font.color.rgb = TEXT_WHITE
    p_ms1.alignment = PP_ALIGN.CENTER
    p_ms2 = tf_ms.add_paragraph()
    p_ms2.text = "Port 8002"
    p_ms2.font.size = Pt(12)
    p_ms2.font.color.rgb = TEXT_WHITE
    p_ms2.alignment = PP_ALIGN.CENTER
    add_footer(slide4, 4)

    # -------------------------------------------------------------
    # SLIDE 5: MICROSERVICES AND REST API ENDPOINTS
    # -------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5, BG_LIGHT)
    add_top_banner(slide5, "Microservices and REST API Endpoints")

    t_shape = slide5.shapes.add_table(4, 4, Inches(0.8), Inches(1.6), Inches(11.7), Inches(4.8))
    table = t_shape.table
    table.columns[0].width = Inches(2.5)
    table.columns[1].width = Inches(1.2)
    table.columns[2].width = Inches(4.5)
    table.columns[3].width = Inches(3.5)

    headers = ["Service", "Port", "Responsibility", "Endpoints"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY_DARK
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.bold = True
        p.font.size = Pt(14)
        p.font.color.rgb = TEXT_WHITE

    rows_data = [
        ("attendance-service", "8003", "Entry point. Performs check-in by calling member and membership services.", "GET /\nPOST /attendance/checkin\nPOST /attendance/checkout/{id}\nGET /attendance/{member_id}"),
        ("member-service", "8001", "Stores and returns gym member profile records.", "GET /\nGET /members\nPOST /members\nGET /members/{id}\nPUT /members/{id}\nDELETE /members/{id}"),
        ("membership-service", "8002", "Stores membership plan tiers and validity status.", "GET /\nGET /memberships\nPOST /memberships\nGET /memberships/{member_id}\nPUT /memberships/{member_id}\nDELETE /memberships/{member_id}")
    ]
    for row_idx, rdata in enumerate(rows_data, start=1):
        for col_idx, cell_value in enumerate(rdata):
            cell = table.cell(row_idx, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = CARD_BG
            p = cell.text_frame.paragraphs[0]
            p.text = cell_value
            p.font.size = Pt(12)
            p.font.color.rgb = TEXT_DARK if col_idx != 0 else NAVY_DARK
            if col_idx == 0:
                p.font.bold = True
    add_footer(slide5, 5)

    # -------------------------------------------------------------
    # SLIDE 6: CHECKPOINT 1 — DEVELOP THE MICROSERVICES (Matching Screenshot Slide 6)
    # -------------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6, BG_LIGHT)
    add_top_banner(slide6, "Checkpoint 1: Develop the Microservices")

    # Left Bullets Box
    tb_c1 = slide6.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    tf_c1 = tb_c1.text_frame
    tf_c1.word_wrap = True

    c1_bullets = [
        "Domain: Gym Management System",
        "Three services, each in its own folder with its own main.py",
        "Built with Python and FastAPI",
        "Each service run and tested independently first",
        "All APIs returned the expected JSON"
    ]
    for b in c1_bullets:
        p = tf_c1.add_paragraph()
        p.text = "•   " + b
        p.font.size = Pt(16)
        p.font.color.rgb = TEXT_DARK
        p.space_before = Pt(18)

    # Right Code Box
    c_code = slide6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.6), Inches(1.6), Inches(5.9), Inches(5.2))
    c_code.fill.solid()
    c_code.fill.fore_color.rgb = NAVY_DARK
    c_code.line.fill.background()

    tb_code = slide6.shapes.add_textbox(Inches(6.8), Inches(1.8), Inches(5.5), Inches(4.8))
    tf_cd = tb_code.text_frame
    tf_cd.word_wrap = True

    code_str = """from fastapi import FastAPI, HTTPException
from .database import Base, engine, SessionLocal
from .models import Member

app = FastAPI(title="Gym Member Service")

Base.metadata.create_all(bind=engine)

@app.get("/members/{member_id}")
def get_member(member_id: int):
    db = SessionLocal()
    member = db.query(Member).filter(
        Member.id == member_id
    ).first()
    db.close()
    if not member:
        raise HTTPException(status_code=404, 
                            detail="Member not found")
    return {"id": member.id, "name": member.name}"""

    p = tf_cd.paragraphs[0]
    p.text = code_str
    p.font.size = Pt(11)
    p.font.name = "Courier New"
    p.font.color.rgb = TEXT_WHITE
    add_footer(slide6, 6)

    # -------------------------------------------------------------
    # SLIDE 7: CHECKPOINT 2 — CONTAINERIZE AND DEPLOY (Matching Screenshot Slide 7)
    # -------------------------------------------------------------
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7, BG_LIGHT)
    add_top_banner(slide7, "Checkpoint 2: Containerize and Deploy")

    # Left Box: Dockerfile
    c_df = slide7.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(2.2), Inches(5.6), Inches(3.6))
    c_df.fill.solid()
    c_df.fill.fore_color.rgb = NAVY_DARK
    c_df.line.fill.background()

    tb_df_head = slide7.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(5.6), Inches(0.5))
    p = tb_df_head.text_frame.paragraphs[0]
    p.text = "Dockerfile (one per service)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = NAVY_DARK

    tb_df_code = slide7.shapes.add_textbox(Inches(0.9), Inches(2.3), Inches(5.4), Inches(3.4))
    tf_df = tb_df_code.text_frame
    tf_df.word_wrap = True
    tf_df.paragraphs[0].text = """FROM python:3.11-slim
WORKDIR /app
COPY requirements.txt .
RUN pip install --no-cache-dir -r requirements.txt
COPY app ./app
EXPOSE 8000
CMD ["uvicorn", "app.main:app", 
     "--host", "0.0.0.0", "--port", "8000"]"""
    tf_df.paragraphs[0].font.size = Pt(11)
    tf_df.paragraphs[0].font.name = "Courier New"
    tf_df.paragraphs[0].font.color.rgb = TEXT_WHITE

    # Right Box: docker-compose.yml
    c_dc = slide7.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.8), Inches(2.2), Inches(5.7), Inches(3.6))
    c_dc.fill.solid()
    c_dc.fill.fore_color.rgb = NAVY_DARK
    c_dc.line.fill.background()

    tb_dc_head = slide7.shapes.add_textbox(Inches(6.8), Inches(1.6), Inches(5.7), Inches(0.5))
    p = tb_dc_head.text_frame.paragraphs[0]
    p.text = "docker-compose.yml (attendance-service)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = NAVY_DARK

    tb_dc_code = slide7.shapes.add_textbox(Inches(6.9), Inches(2.3), Inches(5.5), Inches(3.4))
    tf_dc = tb_dc_code.text_frame
    tf_dc.word_wrap = True
    tf_dc.paragraphs[0].text = """attendance-service:
  build: ./attendance-service
  ports: ["8003:8000"]
  environment:
    - MEMBER_SERVICE_URL=http://member-service:8000
    - MEMBERSHIP_SERVICE_URL=http://membership-service:8000
  networks: [gym-network]"""
    tf_dc.paragraphs[0].font.size = Pt(11)
    tf_dc.paragraphs[0].font.name = "Courier New"
    tf_dc.paragraphs[0].font.color.rgb = TEXT_WHITE

    # Bottom Bullets
    tb_b7 = slide7.shapes.add_textbox(Inches(0.8), Inches(5.9), Inches(11.7), Inches(1.0))
    tf_b7 = tb_b7.text_frame
    b7_list = [
        "Three images built: verified with docker images",
        "One command deploys all: docker compose up --build -d",
        "Three containers running: verified with docker ps"
    ]
    for b in b7_list:
        p = tf_b7.add_paragraph()
        p.text = "•   " + b
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_DARK
        p.space_before = Pt(3)
    add_footer(slide7, 7)

    # -------------------------------------------------------------
    # SLIDE 8: CHECKPOINT 3 — INTER-SERVICE COMMUNICATION (Matching Screenshot Slide 8)
    # -------------------------------------------------------------
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide8, BG_LIGHT)
    add_top_banner(slide8, "Checkpoint 3: Inter-Service Communication")

    # Left White Card Box
    c_dns = slide8.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(3.2))
    c_dns.fill.solid()
    c_dns.fill.fore_color.rgb = CARD_BG
    c_dns.line.color.rgb = RGBColor(226, 232, 240)

    t_bar = slide8.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.6), Inches(0.12), Inches(3.2))
    t_bar.fill.solid()
    t_bar.fill.fore_color.rgb = TEAL_ACCENT
    t_bar.line.fill.background()

    tb_dns = slide8.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(5.1), Inches(2.8))
    tf_dns = tb_dns.text_frame
    tf_dns.word_wrap = True
    p = tf_dns.paragraphs[0]
    p.text = "Docker DNS, not localhost"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = NAVY_DARK

    p = tf_dns.add_paragraph()
    p.text = "Inside a container, localhost means that container itself. Services reach each other by service name on gym-network."
    p.font.size = Pt(13)
    p.font.color.rgb = TEXT_DARK
    p.space_before = Pt(10)

    # Left Bottom Environment Box
    c_env = slide8.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(5.0), Inches(5.6), Inches(1.8))
    c_env.fill.solid()
    c_env.fill.fore_color.rgb = NAVY_DARK
    c_env.line.fill.background()

    tb_env = slide8.shapes.add_textbox(Inches(0.9), Inches(5.1), Inches(5.4), Inches(1.6))
    tf_env = tb_env.text_frame
    tf_env.word_wrap = True
    p = tf_env.paragraphs[0]
    p.text = "MEMBER_SERVICE_URL=http://member-service:8000\nMEMBERSHIP_SERVICE_URL=http://membership-service:8000"
    p.font.size = Pt(11)
    p.font.name = "Courier New"
    p.font.color.rgb = TEXT_WHITE

    # Right Box: End-to-End Request
    tb_req_head = slide8.shapes.add_textbox(Inches(6.8), Inches(1.5), Inches(5.7), Inches(0.5))
    p = tb_req_head.text_frame.paragraphs[0]
    p.text = "End-to-end request"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = NAVY_DARK

    c_req = slide8.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.8), Inches(2.1), Inches(5.7), Inches(4.7))
    c_req.fill.solid()
    c_req.fill.fore_color.rgb = NAVY_DARK
    c_req.line.fill.background()

    tb_req = slide8.shapes.add_textbox(Inches(7.0), Inches(2.3), Inches(5.3), Inches(4.3))
    tf_req = tb_req.text_frame
    tf_req.word_wrap = True
    req_code = """$ curl -X POST http://localhost:8003/attendance/checkin?member_id=1

{
  "message": "Check-in successful",
  "attendance_id": 1,
  "member_id": 1,
  "member_name": "Alice Smith",
  "membership_plan": "VIP_GOLD",
  "membership_status": "ACTIVE",
  "status": "PRESENT",
  "check_in": "2026-10-08 00:20:00"
}"""
    p = tf_req.paragraphs[0]
    p.text = req_code
    p.font.size = Pt(11)
    p.font.name = "Courier New"
    p.font.color.rgb = TEXT_WHITE
    add_footer(slide8, 8)

    # -------------------------------------------------------------
    # SLIDE 9: CHECKPOINT 4 — WORKLOAD GENERATION AND MONITORING (Matching Screenshot Slide 9)
    # -------------------------------------------------------------
    slide9 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide9, BG_LIGHT)
    add_top_banner(slide9, "Checkpoint 4: Workload Generation and Monitoring")

    wl_grid = [
        ("Target API", "POST /attendance/checkin?member_id=1\nTouches all three microservices"),
        ("Load generator", "Python ThreadPoolExecutor / Async benchmark\n30 to 80 requests per level"),
        ("Workload levels", "W1 = 1    W2 = 2    W3 = 4\nW4 = 8    W5 = 16\nconcurrent requests"),
        ("Monitoring", "docker stats sampled during every level; average CPU % and memory MB per container")
    ]
    box_w, box_h, col_s, row_s = Inches(5.6), Inches(2.3), Inches(0.5), Inches(0.3)
    for i, (title, desc) in enumerate(wl_grid):
        col, row = i % 2, i // 2
        left_pos = Inches(0.8) + col * (box_w + col_s)
        top_pos = Inches(1.6) + row * (box_h + row_s)

        box = slide9.shapes.add_shape(MSO_SHAPE.RECTANGLE, left_pos, top_pos, box_w, box_h)
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = RGBColor(226, 232, 240)

        t_bar = slide9.shapes.add_shape(MSO_SHAPE.RECTANGLE, left_pos, top_pos, Inches(0.12), box_h)
        t_bar.fill.solid()
        t_bar.fill.fore_color.rgb = TEAL_ACCENT
        t_bar.line.fill.background()

        tb_b = slide9.shapes.add_textbox(left_pos + Inches(0.25), top_pos + Inches(0.2), box_w - Inches(0.3), box_h - Inches(0.4))
        tf_b = tb_b.text_frame
        tf_b.word_wrap = True

        p_t = tf_b.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(18)
        p_t.font.bold = True
        p_t.font.color.rgb = NAVY_DARK

        p_d = tf_b.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(14)
        p_d.font.color.rgb = TEXT_DARK
        p_d.space_before = Pt(10)

    tb_w_note = slide9.shapes.add_textbox(Inches(0.8), Inches(6.2), Inches(11.7), Inches(0.5))
    p_wn = tb_w_note.text_frame.paragraphs[0]
    p_wn.text = "Recorded for every level: average response time, throughput, failed requests, CPU % and memory MB for all three containers."
    p_wn.font.size = Pt(13)
    p_wn.font.italic = True
    p_wn.font.color.rgb = TEAL_ACCENT
    add_footer(slide9, 9)

    # -------------------------------------------------------------
    # SLIDE 10: OBSERVATION TABLE (MEASURED VALUES) (Matching Screenshot Slide 10)
    # -------------------------------------------------------------
    slide10 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide10, BG_LIGHT)
    add_top_banner(slide10, "Observation Table (measured values)")

    t_shape = slide10.shapes.add_table(6, 11, Inches(0.8), Inches(1.5), Inches(11.7), Inches(4.0))
    t = t_shape.table
    col_widths = [1.0, 1.1, 1.1, 1.2, 0.9, 0.9, 0.9, 0.9, 1.2, 1.2, 1.3]
    for idx, w in enumerate(col_widths):
        t.columns[idx].width = Inches(w)

    h_list = ["Workload", "Concurrency", "Avg resp (ms)", "Throughput (req/s)", "Failed", "CPU % Att", "CPU % Mem", "CPU % MS", "Mem MB Att", "Mem MB Mem", "Mem MB MS"]
    for i, h in enumerate(h_list):
        cell = t.cell(0, i)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY_DARK
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.bold = True
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_WHITE

    obs_data = [
        ("W1", "1", "10.14", "98.38", "0", "0.34%", "0.43%", "0.38%", "60.94", "52.48", "52.43"),
        ("W2", "2", "15.52", "126.71", "0", "0.80%", "0.41%", "0.38%", "66.38", "52.70", "52.91"),
        ("W3", "4", "27.40", "142.50", "0", "0.33%", "0.32%", "0.47%", "67.85", "52.93", "52.90"),
        ("W4", "8", "54.78", "142.29", "0", "0.48%", "0.40%", "0.57%", "85.40", "53.18", "53.14"),
        ("W5", "16", "101.24", "150.50", "0", "0.45%", "0.34%", "0.40%", "90.53", "56.65", "55.61")
    ]
    for r_idx, r_data in enumerate(obs_data, start=1):
        for c_idx, val in enumerate(r_data):
            cell = t.cell(r_idx, c_idx)
            cell.fill.solid()
            # Highlight Attendance Memory
            if c_idx == 8:
                cell.fill.fore_color.rgb = RGBColor(254, 243, 199) # Light amber highlight
            else:
                cell.fill.fore_color.rgb = CARD_BG
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(10)
            p.font.color.rgb = NAVY_DARK if c_idx == 8 else TEXT_DARK
            if c_idx == 0:
                p.font.bold = True

    tb_b10 = slide10.shapes.add_textbox(Inches(0.8), Inches(5.8), Inches(11.7), Inches(1.0))
    tf_b10 = tb_b10.text_frame
    b10_bullets = [
        "0 failed requests at every level",
        "Attendance service (highlighted) uses the most memory (~90 MB); memory stays flat at about 50 to 90 MB"
    ]
    for b in b10_bullets:
        p = tf_b10.add_paragraph()
        p.text = "•   " + b
        p.font.size = Pt(14)
        p.font.color.rgb = TEXT_DARK
        p.space_before = Pt(4)
    add_footer(slide10, 10)

    # -------------------------------------------------------------
    # SLIDE 11: PERFORMANCE GRAPHS
    # -------------------------------------------------------------
    slide11 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide11, BG_LIGHT)
    add_top_banner(slide11, "Performance Graphs")

    chart_files = [
        ("chart_response_time.png", Inches(0.8), Inches(1.5)),
        ("chart_throughput.png", Inches(6.8), Inches(1.5)),
        ("chart_cpu_utilization.png", Inches(0.8), Inches(4.3)),
        ("chart_memory_utilization.png", Inches(6.8), Inches(4.3))
    ]
    for fname, left, top in chart_files:
        if os.path.exists(fname):
            slide11.shapes.add_picture(fname, left, top, width=Inches(5.7), height=Inches(2.6))
    add_footer(slide11, 11)

    # -------------------------------------------------------------
    # SLIDE 12: CONCLUSION AND EVALUATION SUMMARY
    # -------------------------------------------------------------
    slide12 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide12, BG_LIGHT)
    add_top_banner(slide12, "Conclusion and Evaluation Summary")

    c_sum = slide12.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.1))
    c_sum.fill.solid()
    c_sum.fill.fore_color.rgb = CARD_BG
    c_sum.line.color.rgb = RGBColor(226, 232, 240)

    t_bar = slide12.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.6), Inches(0.12), Inches(5.1))
    t_bar.fill.solid()
    t_bar.fill.fore_color.rgb = TEAL_ACCENT
    t_bar.line.fill.background()

    tb_sum = slide12.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(11.1), Inches(4.7))
    tf_sum = tb_sum.text_frame
    tf_sum.word_wrap = True

    p = tf_sum.paragraphs[0]
    p.text = "Summary of Completed Checkpoints"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = NAVY_DARK

    chk_summary = [
        "✔ Checkpoint 1 PASSED: Designed and implemented 3 independent microservices (Member, Membership, Attendance).",
        "✔ Checkpoint 2 PASSED: Containerized each service with Dockerfiles and deployed via Docker Compose.",
        "✔ Checkpoint 3 PASSED: Verified inter-service REST communication over container DNS bridge network.",
        "✔ Checkpoint 4 PASSED: Executed workload testing across 5 concurrency levels (W1 - W5).",
        "✔ Checkpoint 5 PASSED: Analyzed performance measurements, generated visual charts, and verified 0% error rate."
    ]
    for chk in chk_summary:
        p = tf_sum.add_paragraph()
        p.text = chk
        p.font.size = Pt(14)
        p.font.color.rgb = TEXT_DARK
        p.space_before = Pt(14)

    add_footer(slide12, 12)

    # Save presentation
    output_filename = "Gym_Microservices_Lab_Evaluation.pptx"
    prs.save(output_filename)
    print(f"Full 12-slide Presentation generated successfully as '{output_filename}'.")

if __name__ == "__main__":
    create_full_presentation()
