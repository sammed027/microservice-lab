import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette
    DARK_BG = RGBColor(15, 23, 42)      # #0f172a
    CARD_BG = RGBColor(30, 41, 59)      # #1e293b
    ACCENT_BLUE = RGBColor(59, 130, 246) # #3b82f6
    TEXT_WHITE = RGBColor(248, 250, 252) # #f8fafc
    TEXT_MUTED = RGBColor(148, 163, 184) # #94a3b8
    GREEN_ACCENT = RGBColor(16, 185, 129)

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = DARK_BG
        bg.line.fill.background()

    def add_header(slide, title_text, category_text="CLOUD COMPUTING LAB EVALUATION"):
        # Header Box
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.0))
        tf = header_box.text_frame
        tf.word_wrap = True
        
        p_cat = tf.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = ACCENT_BLUE

        p_title = tf.add_paragraph()
        p_title.text = title_text
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_WHITE

    # -------------------------------------------------------------
    # SLIDE 1: TITLE SLIDE
    # -------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1)

    # Accent bar
    bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(2.2), Inches(0.15), Inches(3.2))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT_BLUE
    bar.line.fill.background()

    tb1 = slide1.shapes.add_textbox(Inches(1.2), Inches(2.1), Inches(11.0), Inches(3.5))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "CLOUD COMPUTING LAB EVALUATION"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    p2 = tf1.add_paragraph()
    p2.text = "Build, Deploy and Analyze a Containerized\nMicroservice Application Under Varying Workloads"
    p2.font.size = Pt(28)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_WHITE
    p2.space_before = Pt(10)

    p3 = tf1.add_paragraph()
    p3.text = "Domain: Gym Management System (Member • Membership • Attendance Services)"
    p3.font.size = Pt(16)
    p3.font.color.rgb = GREEN_ACCENT
    p3.space_before = Pt(15)

    # Footer
    tb_foot = slide1.shapes.add_textbox(Inches(1.2), Inches(6.0), Inches(11.0), Inches(0.8))
    tf_foot = tb_foot.text_frame
    p_foot = tf_foot.paragraphs[0]
    p_foot.text = "Technology Stack: Python FastAPI • Docker & Docker Compose • Async HTTP • SQLite • Matplotlib"
    p_foot.font.size = Pt(11)
    p_foot.font.color.rgb = TEXT_MUTED

    # -------------------------------------------------------------
    # SLIDE 2: AIM & MINIMUM ARCHITECTURE
    # -------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2)
    add_header(slide2, "Experiment Aim & System Architecture Flow")

    # Card 1: Aim
    c1 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = ACCENT_BLUE

    tb = slide2.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.2), Inches(4.8))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Project Objectives"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    bullets = [
        "Develop 3 independent RESTful microservices.",
        "Containerize services using Docker Dockerfiles.",
        "Deploy on unified bridge network via Docker Compose.",
        "Establish inter-service communication flow.",
        "Generate 5 varying workload levels (W1 to W5).",
        "Monitor container CPU & Memory utilization.",
        "Perform statistical analysis and graph visualization."
    ]
    for b in bullets:
        p = tf.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(8)

    # Card 2: Architecture Diagram
    c2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    c2.fill.solid()
    c2.fill.fore_color.rgb = CARD_BG
    c2.line.color.rgb = ACCENT_BLUE

    tb_arch = slide2.shapes.add_textbox(Inches(7.0), Inches(1.8), Inches(5.3), Inches(4.8))
    tf_arch = tb_arch.text_frame
    tf_arch.word_wrap = True
    p = tf_arch.paragraphs[0]
    p.text = "Minimum Architecture Flow"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = GREEN_ACCENT

    arch_steps = [
        "Client Request (Port 8003):",
        "   POST /attendance/checkin?member_id=1",
        "",
        "Attendance Service (Orchestrator):",
        "   --> Calls Member Service (Port 8001)",
        "   --> Calls Membership Service (Port 8002)",
        "",
        "Validation Gate:",
        "   Verifies Member Existence & Active Status",
        "",
        "Composite Response:",
        "   Returns Verified Status 'PRESENT' to Client"
    ]
    for step in arch_steps:
        p = tf_arch.add_paragraph()
        p.text = step
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_WHITE if not step.endswith(":") else ACCENT_BLUE

    # -------------------------------------------------------------
    # SLIDE 3: CHECKPOINT 1 — MICROSERVICES DESIGN
    # -------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3)
    add_header(slide3, "Checkpoint 1 — Microservice Responsibilities & REST APIs")

    # Table of Services
    table_shape = slide3.shapes.add_table(4, 4, Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.0))
    table = table_shape.table
    table.columns[0].width = Inches(2.2)
    table.columns[1].width = Inches(3.8)
    table.columns[2].width = Inches(1.2)
    table.columns[3].width = Inches(4.5)

    headers = ["Microservice", "Primary Responsibility", "Port", "Key REST API Endpoints"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.fill.solid()
        cell.fill.fore_color.rgb = ACCENT_BLUE
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.bold = True
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_WHITE

    rows_data = [
        ("Member Service", "Manages gym member profiles (Creation, Retrieval, Updates, Deletions).", "8001", "GET /members\nPOST /members\nPUT /members/{id}\nDELETE /members/{id}"),
        ("Membership Service", "Manages subscription plan tiers (VIP Gold, Silver), validity dates & statuses.", "8002", "GET /memberships\nPOST /memberships\nPUT /memberships/{id}\nDELETE /memberships/{id}"),
        ("Attendance Service", "Orchestrates check-ins/check-outs by querying Member & Membership services.", "8003", "POST /attendance/checkin\nPOST /attendance/checkout/{id}\nGET /attendance/{member_id}")
    ]

    for row_idx, rdata in enumerate(rows_data, start=1):
        for col_idx, cell_value in enumerate(rdata):
            cell = table.cell(row_idx, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = CARD_BG
            p = cell.text_frame.paragraphs[0]
            p.text = cell_value
            p.font.size = Pt(11)
            p.font.color.rgb = TEXT_WHITE

    # -------------------------------------------------------------
    # SLIDE 4: CHECKPOINT 2 — CONTAINERIZATION
    # -------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4)
    add_header(slide4, "Checkpoint 2 — Containerization & Deployment (Docker)")

    # Left Box: Dockerfile
    c_df = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    c_df.fill.solid()
    c_df.fill.fore_color.rgb = CARD_BG
    c_df.line.color.rgb = ACCENT_BLUE

    tb_df = slide4.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.2), Inches(4.8))
    tf_df = tb_df.text_frame
    tf_df.word_wrap = True
    p = tf_df.paragraphs[0]
    p.text = "Microservice Dockerfile Structure"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    df_code = """FROM python:3.11-slim

WORKDIR /app

COPY requirements.txt .
RUN pip install --no-cache-dir -r requirements.txt

COPY app ./app

EXPOSE 8000

CMD ["uvicorn", "app.main:app", 
     "--host", "0.0.0.0", "--port", "8000"]"""
    p = tf_df.add_paragraph()
    p.text = df_code
    p.font.size = Pt(11)
    p.font.name = "Courier New"
    p.font.color.rgb = GREEN_ACCENT
    p.space_before = Pt(10)

    # Right Box: Docker Compose
    c_dc = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    c_dc.fill.solid()
    c_dc.fill.fore_color.rgb = CARD_BG
    c_dc.line.color.rgb = ACCENT_BLUE

    tb_dc = slide4.shapes.add_textbox(Inches(7.0), Inches(1.8), Inches(5.3), Inches(4.8))
    tf_dc = tb_dc.text_frame
    tf_dc.word_wrap = True
    p = tf_dc.paragraphs[0]
    p.text = "Docker Compose Deployment Highlights"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    dc_bullets = [
        "Network Driver: Isolated 'gym-network' (bridge).",
        "Container Ports Mapped: 8001, 8002, 8003.",
        "Persistent Volumes: member-db, membership-db, attendance-db mapped to /app/data/.",
        "Container Inter-Service Dependency: attendance-service depends_on member & membership services.",
        "Environment Environment Variables: Configurable internal URLs (http://member-service:8000)."
    ]
    for b in dc_bullets:
        p = tf_dc.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 5: CHECKPOINT 3 — INTER-SERVICE COMMUNICATION
    # -------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5)
    add_header(slide5, "Checkpoint 3 — Inter-Service Microservice Communication")

    # Sequence Box
    c_seq = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2))
    c_seq.fill.solid()
    c_seq.fill.fore_color.rgb = CARD_BG
    c_seq.line.color.rgb = ACCENT_BLUE

    tb_seq = slide5.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(11.1), Inches(4.8))
    tf_seq = tb_seq.text_frame
    tf_seq.word_wrap = True
    p = tf_seq.paragraphs[0]
    p.text = "Verified End-to-End Inter-Service Request Execution"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    seq_steps = [
        "1. Client triggers HTTP POST http://localhost:8003/attendance/checkin?member_id=1",
        "2. Attendance Service makes HTTP GET to http://member-service:8000/members/1 via Docker DNS.",
        "3. Attendance Service makes HTTP GET to http://membership-service:8000/memberships/1.",
        "4. Attendance Service verifies status == 'ACTIVE' and records timestamped entry into database.",
        "5. Returns HTTP 200 Composite JSON Response:"
    ]
    for s in seq_steps:
        p = tf_seq.add_paragraph()
        p.text = s
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(4)

    json_code = """{
  "message": "Check-in successful",
  "attendance_id": 1,
  "member_id": 1,
  "member_name": "Alice Smith",
  "membership_plan": "VIP_GOLD",
  "membership_status": "ACTIVE",
  "status": "PRESENT",
  "check_in": "2026-10-08 00:20:00"
}"""
    p = tf_seq.add_paragraph()
    p.text = json_code
    p.font.size = Pt(10)
    p.font.name = "Courier New"
    p.font.color.rgb = GREEN_ACCENT
    p.space_before = Pt(6)

    # -------------------------------------------------------------
    # SLIDE 6: CHECKPOINT 4 — WORKLOAD OBSERVATION TABLE
    # -------------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6)
    add_header(slide6, "Checkpoint 4 — Workload Testing & Measured Observation Table")

    # Table
    t_shape = slide6.shapes.add_table(6, 9, Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.0))
    t = t_shape.table
    widths = [0.9, 1.1, 1.0, 0.9, 1.6, 1.5, 1.5, 1.5, 1.7]
    for idx, w in enumerate(widths):
        t.columns[idx].width = Inches(w)

    h_list = ["Test", "Concur.", "Reqs", "Failed", "Avg Latency", "Throughput", "Att Mem", "Mem Mem", "MS Mem"]
    for i, h in enumerate(h_list):
        cell = t.cell(0, i)
        cell.fill.solid()
        cell.fill.fore_color.rgb = ACCENT_BLUE
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_WHITE

    obs_data = [
        ("W1", "1", "30", "0", "10.14 ms", "98.38 req/s", "60.94 MB", "52.48 MB", "52.43 MB"),
        ("W2", "2", "40", "0", "15.52 ms", "126.71 req/s", "66.38 MB", "52.70 MB", "52.91 MB"),
        ("W3", "4", "50", "0", "27.40 ms", "142.50 req/s", "67.85 MB", "52.93 MB", "52.90 MB"),
        ("W4", "8", "60", "0", "54.78 ms", "142.29 req/s", "85.40 MB", "53.18 MB", "53.14 MB"),
        ("W5", "16", "80", "0", "101.24 ms", "150.50 req/s", "90.53 MB", "56.65 MB", "55.61 MB")
    ]
    for r_idx, r_data in enumerate(obs_data, start=1):
        for c_idx, val in enumerate(r_data):
            cell = t.cell(r_idx, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = CARD_BG
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(10)
            p.font.color.rgb = TEXT_WHITE

    # -------------------------------------------------------------
    # SLIDE 7: CHECKPOINT 5 — VISUAL PERFORMANCE GRAPHS
    # -------------------------------------------------------------
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7)
    add_header(slide7, "Checkpoint 5 — Performance Charts & Visual Analysis")

    # Add 4 charts in 2x2 layout
    chart_files = [
        ("chart_response_time.png", Inches(0.8), Inches(1.6)),
        ("chart_throughput.png", Inches(6.8), Inches(1.6)),
        ("chart_cpu_utilization.png", Inches(0.8), Inches(4.5)),
        ("chart_memory_utilization.png", Inches(6.8), Inches(4.5))
    ]

    for fname, left, top in chart_files:
        if os.path.exists(fname):
            slide7.shapes.add_picture(fname, left, top, width=Inches(5.7), height=Inches(2.7))

    # -------------------------------------------------------------
    # SLIDE 8: DEEP-DIVE ANALYSIS & CONCLUSION
    # -------------------------------------------------------------
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide8)
    add_header(slide8, "Statistical Analysis, Evaluator Q&A & Conclusion")

    # Left Box: Key Findings
    c_kf = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    c_kf.fill.solid()
    c_kf.fill.fore_color.rgb = CARD_BG
    c_kf.line.color.rgb = ACCENT_BLUE

    tb_kf = slide8.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.2), Inches(4.8))
    tf_kf = tb_kf.text_frame
    tf_kf.word_wrap = True
    p = tf_kf.paragraphs[0]
    p.text = "Key Performance Findings"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    kf_bullets = [
        "Latency Scaling: Response time increases linearly from 10.14 ms (W1) to 101.24 ms (W5) due to thread queuing and DB write locks.",
        "Throughput Saturation: Peak throughput reached ~150 req/sec at 16 concurrency.",
        "Resource Distribution: Attendance Service consumes highest memory (~90 MB) because it manages async HTTP connection pools to downstream services.",
        "Reliability: 100% success rate (0 failures) achieved across all 5 benchmark levels."
    ]
    for b in kf_bullets:
        p = tf_kf.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(8)

    # Right Box: Conclusion
    c_con = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    c_con.fill.solid()
    c_con.fill.fore_color.rgb = CARD_BG
    c_con.line.color.rgb = ACCENT_BLUE

    tb_con = slide8.shapes.add_textbox(Inches(7.0), Inches(1.8), Inches(5.3), Inches(4.8))
    tf_con = tb_con.text_frame
    tf_con.word_wrap = True
    p = tf_con.paragraphs[0]
    p.text = "Evaluation Summary & Conclusion"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = GREEN_ACCENT

    con_bullets = [
        "Checkpoint 1 PASSED: 3 independent REST microservices.",
        "Checkpoint 2 PASSED: Containerized & deployed with Docker Compose.",
        "Checkpoint 3 PASSED: Inter-service communication verified.",
        "Checkpoint 4 PASSED: Measured workloads W1 - W5.",
        "Checkpoint 5 PASSED: Results analyzed and visualized into graphs.",
        "Project Status: Fully operational & ready for evaluation."
    ]
    for b in con_bullets:
        p = tf_con.add_paragraph()
        p.text = "✔ " + b
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(8)

    # Save presentation
    output_filename = "Gym_Management_Microservices_Presentation.pptx"
    prs.save(output_filename)
    print(f"Presentation saved successfully as '{output_filename}'.")

if __name__ == "__main__":
    build_presentation()
